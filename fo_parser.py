import openpyxl
import json
import sys

def parse_fo_excel(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    result = {}
    
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        
        # Buscar fila de encabezados (primera fila no vacía con múltiples columnas)
        header_row = None
        data_start = 0
        for i, row in enumerate(rows):
            non_empty = [c for c in row if c is not None and str(c).strip()]
            if len(non_empty) >= 2:
                header_row = i
                data_start = i + 1
                break
        
        if header_row is None:
            continue
            
        headers = [str(c).strip() if c is not None else f'Col{j}' for j,c in enumerate(rows[header_row])]
        
        data = []
        for row in rows[data_start:]:
            if all(c is None or str(c).strip() == '' for c in row):
                continue
            record = {}
            for j, cell in enumerate(row):
                if j < len(headers):
                    val = cell
                    if val is not None:
                        record[headers[j]] = str(val).strip()
                    else:
                        record[headers[j]] = ''
            data.append(record)
        
        if data:
            result[name] = data
    
    return result

data = parse_fo_excel('/opt/netquery/data/fo_data.xlsx')
for sheet, rows in data.items():
    print(f"{sheet}: {len(rows)} registros, cols: {list(rows[0].keys()) if rows else []}")

with open('/opt/netquery/data/fo_parsed.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)
print('OK - guardado en fo_parsed.json')
